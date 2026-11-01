#!/usr/bin/env python3
"""Fetch sz159941 (纳指100 ETF) quote and append to daily Excel under data/159941/."""

from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

BEIJING = ZoneInfo("Asia/Shanghai")
REPO_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = REPO_ROOT / "data" / "159941"
QUOTE_URL = "https://qt.gtimg.cn/q=sz159941"

COLUMNS = [
    "记录时间",
    "小时段",
    "现价",
    "昨收",
    "涨跌额",
    "涨跌幅(%)",
    "开盘价",
    "行情时间",
]


def is_trading_moment(now: datetime) -> bool:
    try:
        import chinese_calendar as cc
    except ImportError:
        print("chinese_calendar not installed", file=sys.stderr)
        return False

    d = now.date()
    if not cc.is_workday(d):
        return False
    t = now.time()
    from datetime import time

    morning = time(9, 30) <= t <= time(11, 59)
    afternoon = time(13, 0) <= t <= time(15, 59)
    return morning or afternoon


def fetch_quote() -> dict[str, str | float]:
    req = Request(QUOTE_URL, headers={"User-Agent": "Mozilla/5.0"})
    raw = urlopen(req, timeout=15).read()
    text = raw.decode("gbk", errors="replace")
    m = re.search(r'v_sz159941="([^"]+)"', text)
    if not m:
        raise RuntimeError("unexpected quote response")
    parts = m.group(1).split("~")
    if len(parts) < 33:
        raise RuntimeError(f"too few fields: {len(parts)}")

    def f(idx: int) -> float:
        try:
            return float(parts[idx])
        except (ValueError, IndexError):
            return float("nan")

    quote_time = parts[30] if len(parts) > 30 else ""
    return {
        "name": parts[1],
        "code": parts[2],
        "price": f(3),
        "prev_close": f(4),
        "open": f(5),
        "change": f(31),
        "change_pct": f(32),
        "quote_time": quote_time,
    }


def hour_slot(now: datetime) -> str:
    return now.strftime("%Y-%m-%d %H:00")


def load_workbook_path(day: str) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    return DATA_DIR / f"{day}.xlsx"


def append_or_update_row(path: Path, row: dict[str, str | float], slot: str) -> None:
    from openpyxl import Workbook, load_workbook

    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "159941"
        ws.append(COLUMNS)

    header = [cell.value for cell in ws[1]]
    if header != COLUMNS:
        ws.delete_rows(1, ws.max_row)
        ws.append(COLUMNS)

    slot_col = COLUMNS.index("小时段") + 1
    target_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=slot_col).value == slot:
            target_row = r
            break

    values = [row[c] for c in COLUMNS]
    if target_row is not None:
        for col, val in enumerate(values, start=1):
            ws.cell(row=target_row, column=col, value=val)
    else:
        ws.append(values)

    wb.save(path)


def main() -> int:
    now = datetime.now(BEIJING)
    if not is_trading_moment(now):
        print(f"skip: outside trading window ({now.isoformat()})")
        return 0

    q = fetch_quote()
    day = now.strftime("%Y-%m-%d")
    slot = hour_slot(now)
    row = {
        "记录时间": now.strftime("%Y-%m-%d %H:%M:%S"),
        "小时段": slot,
        "现价": q["price"],
        "昨收": q["prev_close"],
        "涨跌额": q["change"],
        "涨跌幅(%)": q["change_pct"],
        "开盘价": q["open"],
        "行情时间": q["quote_time"],
    }
    path = load_workbook_path(day)
    append_or_update_row(path, row, slot)
    print(f"updated {path} slot={slot} pct={q['change_pct']}%")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
